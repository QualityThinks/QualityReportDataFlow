import os
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import mysql.connector
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from dotenv import load_dotenv

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

SQL_QUERY = """
SELECT * FROM `process_result`
"""

OUTPUT_DIR = BASE_DIR / "output"
SHEET_NAME = "process_result"
TABLE_NAME = "ProcessResult"


def require_env(name, allow_empty=False):
    value = os.environ.get(name)

    if value is None or (not allow_empty and not value.strip()):
        raise RuntimeError(
            f"{name} environment variable has not been configured. "
            f"Add it to {BASE_DIR / '.env'}"
        )

    return value.strip()


MYSQL_HOST = require_env("MYSQL_HOST")
MYSQL_PORT = int(require_env("MYSQL_PORT"))
MYSQL_DATABASE = require_env("MYSQL_DATABASE")
MYSQL_USER = require_env("MYSQL_USER")
MYSQL_PASSWORD = require_env("MYSQL_PASSWORD", allow_empty=True)

# Email configuration
EMAIL_RECIPIENT = require_env("EMAIL_RECIPIENT")
EMAIL_SENDER = require_env("EMAIL_SENDER")
EMAIL_SUBJECT = os.environ.get("EMAIL_SUBJECT", "Quality Report").strip()
EMAIL_BODY = os.environ.get(
    "EMAIL_BODY",
    "Hello,\n\nPlease find the attached quality report.\n\nRegards,\nAutomation",
).strip()


def connect_to_mysql(host, port, database, user, password):
    return mysql.connector.connect(
        host=host,
        port=port,
        database=database,
        user=user,
        password=password
    )


def _create_table(worksheet, table_name, num_columns, num_rows):
    last_col = get_column_letter(num_columns)
    last_row = num_rows + 1  # +1 for the header row
    table = Table(displayName=table_name, ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def write_xlsx(dataframe, xlsx_path):
    """Create a fresh xlsx with the data inside an Excel table."""
    columns = list(dataframe.columns)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME

    columns = list(dataframe.columns)
    worksheet.append(columns)
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append(list(row))

    _create_table(worksheet, TABLE_NAME, len(columns), len(dataframe))
    workbook.save(xlsx_path)


def write_eml(xlsx_path, eml_path):
    """Build an .eml email file with the xlsx attached."""
    message = EmailMessage()
    message["From"] = EMAIL_SENDER
    message["To"] = EMAIL_RECIPIENT
    message["Subject"] = EMAIL_SUBJECT
    message["Date"] = datetime.now().astimezone().strftime("%a, %d %b %Y %H:%M:%S %z")
    message.set_content(EMAIL_BODY)

    with open(xlsx_path, "rb") as attachment:
        message.add_attachment(
            attachment.read(),
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=os.path.basename(xlsx_path),
        )

    with open(eml_path, "wb") as eml_file:
        eml_file.write(bytes(message))


def export_query_to_email(connection, output_dir):

    try:
        print("Connected to MySQL successfully.")

        dataframe = pd.read_sql_query(
            SQL_QUERY,
            connection
        )

        if dataframe.empty:
            print("Query returned no rows. Nothing to export.")
            return

        output_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xlsx_path = output_dir / f"result_{timestamp}.xlsx"
        eml_path = output_dir / f"result_{timestamp}.eml"

        write_xlsx(dataframe, xlsx_path)
        write_eml(xlsx_path, eml_path)

        print("Export completed successfully.")
        print(f"Rows exported: {len(dataframe)}")
        print(f"Excel file: {xlsx_path}")
        print(f"Email file: {eml_path} (to: {EMAIL_RECIPIENT})")

    except mysql.connector.Error as error:
        print(f"MySQL error: {error}")

    except PermissionError:
        print(
            "Permission denied. Make sure the output files are not open "
            "and you have access to the output folder."
        )

    except Exception as error:
        print(f"Unexpected error: {error}")

    finally:
        if connection is not None and connection.is_connected():
            connection.close()
            print("MySQL connection closed.")


if __name__ == "__main__":
    connection = connect_to_mysql(
        MYSQL_HOST,
        MYSQL_PORT,
        MYSQL_DATABASE,
        MYSQL_USER,
        MYSQL_PASSWORD
    )
    export_query_to_email(connection, OUTPUT_DIR)
